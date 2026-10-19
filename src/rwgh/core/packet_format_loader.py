"""数据包格式 xlsx -> 报文 hex 构造器。

「任务规划输入输出数据包」xlsx 的每个 sheet 是一张逐字段规格表:
  第 1 行表头: 序号 | 字段 | (子字段) | 字节数 | 数据类型 | 备注 | 测试数据 | 十六进制码
  第 2 行起每行一个字段,「十六进制码」列是该字段按「字节数」的大端十六进制。

简单 sheet(星时/轨道数据/轨道根数): 依次拼接各字段的「十六进制码」即为完整报文。
记录循环 sheet(数传链路资源/数传站址资源/固存状态信息/常规管控成像任务_点目标):
  记录字段在多列上横向重复(每条记录一组「测试数据/十六进制码」列对)。
  拼接顺序 = 顶层字段(只有一组值的行) + 按记录序号 k=1..N 依次输出第 k 组的各行字段。
  记录条数 N 取自对应的「数量」字段(其十六进制码值)。

sheet -> 报文标识: 星时→0B01, 轨道数据→A002, 轨道根数→A012,
常规管控成像任务_点目标→3101, 数传链路资源→0C12,
数传站址资源→0C15, 固存状态信息→0802。
"""
from __future__ import annotations

import logging
import re
from typing import Optional

from openpyxl import load_workbook

_log = logging.getLogger("rwgh.core.packet_format_loader")

# sheet 名 -> 报文标识
SHEET_TO_CODE: dict[str, str] = {
    "星时": "0B01",
    "轨道数据": "A002",
    "轨道根数": "A012",
    "常规管控成像任务_点目标": "3101",
    "数传链路资源": "0C12",
    "数传站址资源": "0C15",
    "固存状态信息": "0802",
}

# 表头识别关键字
_HEADER_NAMES = ("序号", "字段", "子字段", "字节数", "数据类型", "注释",
                 "测试数据", "十六进制码")

_HEX_RE = re.compile(r"^[0-9A-Fa-f]+$")


class PacketFormatError(Exception):
    """数据包格式 xlsx 解析错误。"""


def _cell_str(ws, row: int, col: int) -> str:
    v = ws.cell(row=row, column=col).value
    return str(v).strip() if v is not None else ""


def _hex_val(ws, row: int, col: int,
             nbytes: Optional[int] = None) -> Optional[str]:
    """读取单元格并返回合法十六进制串(偶数长度),非法返回 None。

    nbytes 为该字段的「字节数」:若单元格 hex 长于 nbytes,按右对齐取末尾
    nbytes 字节(用于修正个别 sheet 中符号扩展多写 FF/00 的写法);
    短于 nbytes 时左侧补零。"""
    s = _cell_str(ws, row, col).replace(" ", "").replace("\n", "")
    if not s or len(s) % 2 != 0:
        return None
    if not _HEX_RE.match(s):
        return None
    s = s.upper()
    if nbytes is not None and nbytes > 0:
        want = nbytes * 2
        if len(s) > want:
            s = s[-want:]
        elif len(s) < want:
            s = s.rjust(want, "0")
    return s


def _find_header_cols(ws) -> dict[str, int]:
    """扫描第 1 行表头,返回 {列名: 列号}(各取首次出现)。"""
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = _cell_str(ws, 1, c)
        if h in _HEADER_NAMES and h not in cols:
            cols[h] = c
    return cols


def build_packet_hex_diag(ws) -> tuple[bytes, list[str], list[str]]:
    """按 sheet 的逐字段规格表构造完整报文 hex 字节,并返回分级校验诊断。

    返回 (bytes, hard_warns, soft_warns):
    - hard: 影响数据正确性的异常(字段缺hex/记录不齐/错位),UI 需弹窗告警
    - soft: 非致命提示(如声明包长与实际不符),状态栏提示即可

    致命问题(无法构造出报文)抛 PacketFormatError。
    ws: openpyxl worksheet。
    """
    cols = _find_header_cols(ws)
    c_hex0 = cols.get("十六进制码")
    c_field = cols.get("字段")
    c_bytes = cols.get("字节数")
    if c_hex0 is None:
        raise PacketFormatError("未找到「十六进制码」表头列")
    hard: list[str] = []
    soft: list[str] = []

    def row_nbytes(r: int) -> Optional[int]:
        if c_bytes is None:
            return None
        v = ws.cell(row=r, column=c_bytes).value
        try:
            return int(v)
        except (TypeError, ValueError):
            return None

    def row_name(r: int) -> str:
        if c_field is None:
            return ""
        n = _cell_str(ws, r, c_field)
        return n

    def field_hex(r: int, col: int) -> Optional[str]:
        return _hex_val(ws, r, col, row_nbytes(r))

    # ---- 探测记录循环「数量」字段行(排除「总包数量」,取最后出现者)----
    count_row: Optional[int] = None
    count_val: Optional[int] = None
    if c_field is not None:
        for r in range(2, ws.max_row + 1):
            name = _cell_str(ws, r, c_field)
            if "数量" in name and name != "总包数量":
                hv = field_hex(r, c_hex0)
                if hv is not None:
                    count_row = r
                    count_val = int(hv, 16)
                else:
                    raise PacketFormatError(
                        f"记录数量字段[{name}]缺少有效十六进制码,"
                        "无法确定循环记录条数")

    parts: list[str] = []

    if count_row is None or count_val is None:
        # 简单 sheet: 逐行拼接第 1 组「十六进制码」
        for r in range(2, ws.max_row + 1):
            hv = field_hex(r, c_hex0)
            if hv is not None:
                parts.append(hv)
            elif row_name(r) and row_nbytes(r):
                hard.append(f"字段[{row_name(r)}]缺少有效十六进制码,已跳过"
                            f"({row_nbytes(r)}B,报文将相应缩短)")
    else:
        # 顶层字段: 第 2 行 ~ 数量字段行
        for r in range(2, count_row + 1):
            hv = field_hex(r, c_hex0)
            if hv is not None:
                parts.append(hv)
            elif row_name(r) and row_nbytes(r):
                hard.append(f"顶层字段[{row_name(r)}]缺少有效十六进制码,"
                            f"已跳过({row_nbytes(r)}B)")
        # 记录字段: 按记录序号 k 依次取第 k 组(列对步长 2)
        # 记录布局边界: 从数量行下一行起(首行「字段」列可为「链路1/目标1信息」
        # 等记录名),遇到「字段」列再次出现非空值(如「……」)即认为记录布局
        # 结束;其后是布局说明行(「链路K/目标M信息」,有字节数无hex)与
        # 尾部顶层字段(如0C12尾部「预留4B」,仅第1组有hex,只拼一次)。
        rec_rows: list[int] = []
        r = count_row + 1
        while r <= ws.max_row:
            name2 = _cell_str(ws, r, c_field) if c_field is not None else ""
            if name2 and rec_rows:
                break
            rec_rows.append(r)
            r += 1
        layout_len = sum(nb for r2 in rec_rows
                         if (nb := row_nbytes(r2)) and nb > 0)
        per_rec: list[int] = []
        for k in range(1, count_val + 1):
            hc = c_hex0 + 2 * (k - 1)
            got = 0
            for r2 in rec_rows:
                hv = field_hex(r2, hc)
                if hv is not None:
                    parts.append(hv)
                    got += len(hv) // 2
            per_rec.append(got)
        # 记录完整性: 各条记录实取字节数应一致且等于布局定义长度
        if per_rec and len(set(per_rec)) > 1:
            bad = [f"第{k+1}条{v}B" for k, v in enumerate(per_rec)
                   if v != per_rec[0]]
            hard.append("记录间字节数不一致(第1条"
                        f"{per_rec[0]}B,{'、'.join(bad)}),报文已错位")
        elif per_rec and layout_len and per_rec[0] != layout_len:
            hard.append(f"记录hex不全(布局{layout_len}B/实取{per_rec[0]}B),"
                        "报文将错位缩短")
        # 记录区之后的尾部顶层字段(有字节数+第1组有hex才拼入;
        # 「链路K/目标M信息」等说明行有字节数无hex,静默跳过;
        # 真实尾部字段缺hex导致的长度短缺由下方「声明包长」软校验兜底)
        for r2 in range(r, ws.max_row + 1):
            nb = row_nbytes(r2)
            if not nb or nb <= 0:
                continue
            name2 = _cell_str(ws, r2, c_field) if c_field is not None else ""
            sub2 = (_cell_str(ws, r2, c_field + 1)
                    if c_field is not None else "")
            if not (name2 or sub2):
                continue
            hv = field_hex(r2, c_hex0)
            if hv is not None:
                parts.append(hv)

    # ---- 声明包长(「当前包长度」字段)与实际构造长度交叉校验(软提示) ----
    built_len = sum(len(p) for p in parts) // 2
    if built_len and c_field is not None:
        for r in range(2, ws.max_row + 1):
            if _cell_str(ws, r, c_field) == "当前包长度":
                hv = field_hex(r, c_hex0)
                if hv is not None and int(hv, 16) != built_len:
                    soft.append(f"声明包长{int(hv, 16)}≠实际构造{built_len}B"
                                "(以实际构造为准)")
                break

    hex_str = "".join(parts)
    if not hex_str:
        raise PacketFormatError("sheet 未构造出任何有效字段")
    try:
        return bytes.fromhex(hex_str), hard, soft
    except ValueError as exc:
        raise PacketFormatError(f"构造的 hex 非法: {exc}") from exc


def build_packet_hex(ws) -> bytes:
    """按 sheet 构造报文 hex(兼容接口,不含诊断)。"""
    data, _, _ = build_packet_hex_diag(ws)
    return data


def load_packets_with_diag(path: str) -> tuple[dict[str, bytes], list[str], list[str]]:
    """加载 xlsx,对已知 sheet 构造报文 hex,并汇总分级校验诊断。

    返回 (packets, hard_warns, soft_warns):
    - hard: 影响数据正确性(缺sheet/构造失败/缺hex/记录不齐),UI 需弹窗
    - soft: 非致命提示(声明包长不符),状态栏提示即可
    """
    wb = load_workbook(path, data_only=True)
    try:
        out: dict[str, bytes] = {}
        hard: list[str] = []
        soft: list[str] = []
        for sheet, code in SHEET_TO_CODE.items():
            if sheet not in wb.sheetnames:
                hard.append(f"{code}: 缺少sheet[{sheet}],该报文未加载")
                _log.warning("xlsx 缺少 sheet: %s", sheet)
                continue
            try:
                data, wh, ws_ = build_packet_hex_diag(wb[sheet])
                out[code] = data
                hard.extend(f"{code}: {x}" for x in wh)
                soft.extend(f"{code}: {x}" for x in ws_)
            except PacketFormatError as exc:
                hard.append(f"{code}: 构造失败({exc})")
                _log.warning("sheet %s 构造失败: %s", sheet, exc)
        return out, hard, soft
    finally:
        wb.close()


def load_packets_from_xlsx(path: str) -> dict[str, bytes]:
    """加载 xlsx,对已知 sheet 构造报文 hex(兼容接口,不含诊断)。

    返回 {"0B01": b"...", "3101": b"...", ...};未包含的 sheet 跳过。
    """
    wb = load_workbook(path, data_only=True)
    try:
        out: dict[str, bytes] = {}
        for sheet, code in SHEET_TO_CODE.items():
            if sheet not in wb.sheetnames:
                _log.warning("xlsx 缺少 sheet: %s", sheet)
                continue
            try:
                out[code] = build_packet_hex(wb[sheet])
            except PacketFormatError as exc:
                _log.warning("sheet %s 构造失败: %s", sheet, exc)
        return out
    finally:
        wb.close()
