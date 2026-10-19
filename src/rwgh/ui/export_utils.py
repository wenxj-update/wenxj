"""数据导出工具:Excel / TXT / JSON,带文件对话框。

供各视图的 export_data() 结果统一导出,所有方法失败返回 False。
"""
from __future__ import annotations

import json

from PySide6.QtWidgets import QFileDialog


class ExportUtils:
    """静态导出工具集。"""

    @staticmethod
    def export_excel(rows: list[dict], columns: list[str],
                    filepath: str, sheet_name: str = "Sheet1") -> bool:
        """导出到 Excel (xlsx),用 openpyxl。

        第一行写表头(columns),后续行写数据(dict.get(col, ""))。
        """
        try:
            import openpyxl
        except ImportError:
            return False
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(list(columns))
            for row in rows:
                ws.append([row.get(col, "") for col in columns])
            wb.save(filepath)
            return True
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def export_excel_multi(filepath: str,
                           sheets: list[tuple[str, list[dict]]]) -> bool:
        """多 Sheet 导出:每个 tuple(sheet 名, 行列表)占一个 sheet。"""
        try:
            import openpyxl
        except ImportError:
            return False
        try:
            wb = openpyxl.Workbook()
            # 删除默认 sheet
            wb.remove(wb.active)
            for sheet_name, rows in sheets:
                # sheet 名最多 31 字符
                ws = wb.create_sheet(title=sheet_name[:31])
                if not rows:
                    ws.append(["(空)"])
                    continue
                columns = list(rows[0].keys())
                ws.append(list(columns))
                for row in rows:
                    ws.append([row.get(col, "") for col in columns])
            wb.save(filepath)
            return True
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def export_txt(rows: list[dict], columns: list[str],
                   filepath: str) -> bool:
        """导出到 TXT (制表符分隔)。

        第一行 tab-join(columns),后续行 tab-join(str(val))。
        """
        try:
            lines = ["\t".join(str(c) for c in columns)]
            for row in rows:
                lines.append(
                    "\t".join(str(row.get(col, "")) for col in columns))
            with open(filepath, "w", encoding="utf-8") as f:
                f.write("\n".join(lines))
            return True
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def export_json(rows: list[dict], filepath: str) -> bool:
        """导出到 JSON(ensure_ascii=False,缩进 2,默认 str 序列化)。"""
        try:
            text = json.dumps(rows, ensure_ascii=False, indent=2,
                              default=str)
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(text)
            return True
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def export_dialog(parent, rows: list[dict], columns: list[str],
                      default_name: str = "export") -> bool:
        """弹出 QFileDialog 让用户选格式(xlsx/txt/json),然后导出。

        根据所选文件后缀调用对应方法;用户取消返回 False。
        """
        try:
            path, _ = QFileDialog.getSaveFileName(
                parent, "导出数据", default_name,
                "Excel (*.xlsx);;文本 (*.txt);;JSON (*.json)")
            if not path:
                return False
            low = path.lower()
            if low.endswith(".xlsx"):
                return ExportUtils.export_excel(rows, columns, path)
            if low.endswith(".txt"):
                return ExportUtils.export_txt(rows, columns, path)
            if low.endswith(".json"):
                return ExportUtils.export_json(rows, path)
            return False
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def export_all_dialog(parent,
                          pages: list[tuple[str, list[dict]]],
                          default_name: str = "全部导出") -> bool:
        """全部导出:每个页面一个 sheet。

        pages: [(页面名, 行列表), ...]
        """
        try:
            path, _ = QFileDialog.getSaveFileName(
                parent, "全部导出", default_name,
                "Excel (*.xlsx)")
            if not path:
                return False
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            return ExportUtils.export_excel_multi(path, pages)
        except Exception:  # noqa: BLE001
            return False
